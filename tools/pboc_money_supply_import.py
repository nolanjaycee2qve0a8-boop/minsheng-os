"""Read-only PBOC money-supply workbook staging tool.

It never changes an official workbook and never writes Minsheng OS data.  It
prints JSON (or writes an explicitly requested staging JSON) for review.
"""
from __future__ import annotations
import argparse, glob, hashlib, importlib.util, json, os, re
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
PERIOD=re.compile(r"^(20\d{2})[./-](\d{1,2})$")
def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()
def period(value):
    match=PERIOD.match(clean(value))
    return f"{match.group(1)}-{int(match.group(2)):02d}" if match else None
def clean(value): return re.sub(r"\s|\u00a0", "", str(value or ''))
def numeric(value):
    try: return float(clean(value).replace(',',''))
    except (TypeError,ValueError): return None
def workbook_base(path):
    return {'fileName':path.name,'relativePath':str(path.relative_to(ROOT)).replace('\\','/'),'sha256':digest(path),'sourceId':'pboc','format':path.suffix.lower()[1:],'year':None}
def xlsx(path):
    from openpyxl import load_workbook
    result=workbook_base(path); result.update({'reader':'openpyxl','readerStatus':'READ_ONLY_OK','sheets':[],'records':[]})
    book=load_workbook(path,read_only=True,data_only=True)
    for ws in book.worksheets:
        rows=list(ws.iter_rows(values_only=True)); headers=[]; m2row=None; unit=None
        header_rows=[]
        for ri,row in enumerate(rows,1):
            for ci,value in enumerate(row,1):
                text=str(value or '')
                if not unit and '单位' in text and '亿元' in text: unit='亿元人民币'
                if period(value): header_rows.append((ri,ci,period(value),text))
            if m2row is None and any('货币和准货币' in clean(v) and 'M2' in clean(v) for v in row): m2row=ri
        raw_headers=[(ci,p,text) for ri,ci,p,text in header_rows if ri==max((r for r,_,_,_ in header_rows if r<m2row),default=0)] if m2row else []
        # Official workbooks display October as e.g. "2023.1" in one column,
        # between September and November. Resolve only that unambiguous sequence.
        for i,(ci,p,text) in enumerate(raw_headers):
            prev=raw_headers[i-1][1] if i else None; nxt=raw_headers[i+1][1] if i+1<len(raw_headers) else None
            if p.endswith('-01') and prev and nxt and prev.endswith('-09') and nxt.endswith('-11'): headers.append((ci,p[:5]+'10'))
            else: headers.append((ci,p))
        sheet={'name':ws.title,'rows':ws.max_row,'columns':ws.max_column,'detectedDateFields':[p for _,p in headers],'detectedM2Fields':['货币和准货币（M2）'] if m2row else [],'detectedUnits':[unit] if unit else []}
        result['sheets'].append(sheet)
        if m2row and headers:
            values=rows[m2row-1]
            for ci,p in headers:
                value=values[ci-1] if ci<=len(values) else None
                amount=numeric(value)
                if amount is not None:
                    result['records'].append({'period':p,'value':amount,'unit':unit,'sheet':ws.title,'row':m2row,'column':ci,'originalValue':value})
    years=sorted({int(r['period'][:4]) for r in result['records']}); result['year']=years[0] if len(years)==1 else None
    return result
def xls(path):
    result=workbook_base(path)
    if not importlib.util.find_spec('xlrd'):
        result.update({'reader':'xlrd','readerStatus':'XLS_READER_UNAVAILABLE','sheets':[],'records':[]}); return result
    import xlrd
    book=xlrd.open_workbook(path,on_demand=True); result.update({'reader':'xlrd','readerVersion':xlrd.__version__,'readerStatus':'READ_ONLY_OK','sheets':[],'records':[]})
    for name in book.sheet_names():
        ws=book.sheet_by_name(name); rows=[ws.row_values(i) for i in range(ws.nrows)]; header_rows=[]; m2row=None; unit=None
        for ri,row in enumerate(rows,1):
            for ci,value in enumerate(row,1):
                text=str(value or '')
                if not unit and '单位' in text and '亿元' in text: unit='亿元人民币'
                if period(value): header_rows.append((ri,ci,period(value),text))
            if m2row is None and any('货币和准货币' in clean(v) or 'Money&Quasi-money' in clean(v) for v in row): m2row=ri
        raw_headers=[(ci,p,text) for ri,ci,p,text in header_rows if ri==max((r for r,_,_,_ in header_rows if r<m2row),default=0)] if m2row else []
        headers=[]
        for i,(ci,p,text) in enumerate(raw_headers):
            prev=raw_headers[i-1][1] if i else None; nxt=raw_headers[i+1][1] if i+1<len(raw_headers) else None
            if p.endswith('-01') and prev and nxt and prev.endswith('-09') and nxt.endswith('-11'): headers.append((ci,p[:5]+'10'))
            else: headers.append((ci,p))
        result['sheets'].append({'name':name,'rows':ws.nrows,'columns':ws.ncols,'detectedDateFields':[p for _,p in headers],'detectedM2Fields':['M2'] if m2row else [],'detectedUnits':[unit] if unit else []})
        if m2row and headers:
            values=rows[m2row-1]
            for ci,p in headers:
                value=values[ci-1] if ci<=len(values) else None
                amount=numeric(value)
                if amount is not None:
                    result['records'].append({'period':p,'value':amount,'unit':unit,'sheet':name,'row':m2row,'column':ci,'originalValue':value})
    years=sorted({int(r['period'][:4]) for r in result['records']}); result['year']=years[0] if len(years)==1 else None
    return result
def main():
    parser=argparse.ArgumentParser(); parser.add_argument('--output'); args=parser.parse_args()
    paths=[]
    for folder in ('sources/manul','sources/manual'):
        paths.extend(Path(p) for p in glob.glob(str(ROOT/folder/'**/*'),recursive=True) if Path(p).suffix.lower() in {'.xls','.xlsx'})
    books=[xlsx(p) if p.suffix.lower()=='.xlsx' else xls(p) for p in sorted(paths)]
    payload={'tool':'pboc_money_supply_import','mode':'READ_ONLY_STAGING','workbooks':books,'balanceRecords':[r|{'fileName':b['fileName'],'sha256':b['sha256'],'sourceId':'pboc'} for b in books for r in b['records']]}
    text=json.dumps(payload,ensure_ascii=False,indent=2)
    if args.output: Path(args.output).write_text(text,encoding='utf-8')
    else: print(text)
if __name__=='__main__': main()
